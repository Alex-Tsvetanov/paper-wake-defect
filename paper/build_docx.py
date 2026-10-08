r"""Build the TELECOM 2026 Word copy of P1 from its LaTeX sources.

One command, from any directory:

    python paper/build_docx.py               # writes paper/Tsvetanov-wake-TELECOM2026.docx
    python paper/build_docx.py --word-check  # also opens it in Word (Windows): page count + PDF

Re-run it after any edit to main.tex, refs.bib, fig-*.tex or results/*. The steps:
 1. latexmk builds main.tex into paper/build-docx/latex (BIBINPUTS points at paper/), so
    main.aux and main.bbl are fresh and a parallel build of paper/main.pdf is not disturbed.
 2. Each figure's \input{fig-...} is compiled standalone with main.tex's own preamble (same
    class, column width and fonts) and rendered to PNG at FIG_DPI.
 3. main.tex is read by a strict reader for the TeX subset the paper uses. Macros from
    ../results/rework/macros.tex and main.tex are expanded, so no number is retyped. \ref and \cite
    numbers come from main.aux; the reference list comes from main.bbl, in the LaTeX order.
 4. The body of template.docx (TELECOM 2026 IEEE template) is replaced. Its styles number the
    headings, figure captions and references, so none of those numbers is typed here.
 5. The .docx is written with fixed zip timestamps: the same sources give the same bytes.
An unknown command or environment stops the build with file:line. No en or em dash is written:
LaTeX -- and --- become "-", and the Abstract/Keywords labels end in LABEL_SEP.
Needs python-docx, PyMuPDF and latexmk with pdflatex (TeX Live or MiKTeX).
"""
import argparse
import copy
import os
import re
import subprocess
import sys
import unicodedata
import zipfile
from collections import namedtuple
from xml.sax.saxutils import escape as xml_escape

import docx
import fitz  # PyMuPDF
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor, Twips

HERE = os.path.dirname(os.path.abspath(__file__))
MAIN_TEX = os.path.join(HERE, "main.tex")
TEMPLATE = os.path.join(HERE, "template.docx")
OUTPUT = os.path.join(HERE, "Tsvetanov-wake-TELECOM2026.docx")
BUILD = os.path.join(HERE, "build-docx")
LATEX_OUT = os.path.join(BUILD, "latex")
FIG_OUT = os.path.join(BUILD, "fig")

FIG_DPI = 600
MAX_PAGES = 4
LABEL_SEP = ":"                 # the template writes "Abstract" + em dash; the author allows no dashes
COLUMN_TWIPS = 4866             # body column of the template: (11906 - 2*907 - 360) / 2
NBSP, THIN, ZWSP, MEDSP = "\u00a0", "\u202f", "\u200b", "\u2005"
MATH_FONT, CODE_FONT, TEXT_FONT = "Cambria Math", "Courier New", "Times New Roman"
NOT_IN_TIMES = set("\u2208\u2209\u220b\u2282\u2286\u2283\u2287\u222a\u2229\u2200\u2203\u2207\u2217")
FORBIDDEN = {"\u2013": "en dash", "\u2014": "em dash"}
COLORS = {"red": "FF0000", "blue": "0000FF", "black": "000000", "green": "008000"}


class TexError(Exception):
    pass


def read(path):
    with open(path, encoding="utf-8") as f:
        return f.read()


# ------------------------------------------------------------------ tokenizer
Tok = namedtuple("Tok", "kind val where")  # cs char space par bg eg math tie param


def tokenize(src, fname):
    """TeX's reading rules: a comment eats its newline, spaces after a control word are
    skipped, a blank line is a paragraph break, other newlines are spaces."""
    toks, state, line, i, n = [], "N", 1, 0, len(src)

    def add(kind, val=None):
        toks.append(Tok(kind, val, f"{fname}:{line}"))

    while i < n:
        c = src[i]
        if c == "%":
            j = src.find("\n", i)
            i, line, state = (n if j < 0 else j + 1), line + 1, "N"
        elif c == "\n":
            if state == "N":
                add("par")
            elif state == "M":
                add("space")
            i, line, state = i + 1, line + 1, "N"
        elif c in " \t\r":
            if state == "M":
                add("space")
                state = "S"
            i += 1
        elif c == "\\":
            j = i + 1
            if j < n and src[j].isalpha():
                while j < n and src[j].isalpha():
                    j += 1
                add("cs", src[i + 1:j])
                i, state = j, "S"
            elif j < n:
                add("cs", " " if src[j] == "\n" else src[j])
                line += src[j] == "\n"
                i, state = j + 1, ("S" if src[j] in " \n" else "M")
            else:
                raise TexError(f"{fname}:{line}: backslash at end of file")
        elif c == "#" and i + 1 < n and src[i + 1].isdigit():
            add("param", int(src[i + 1]))
            i, state = i + 2, "M"
        else:
            add({"{": "bg", "}": "eg", "$": "math", "~": "tie"}.get(c, "char"), c)
            i, state = i + 1, "M"
    return toks


def tok_text(t):
    if t.kind in ("char", "cs") and (t.kind == "char" or len(t.val) == 1):
        return " " if t.val == " " else t.val
    return {"space": " ", "tie": "~", "bg": "{", "eg": "}", "math": "$"}.get(t.kind, "\\" + str(t.val))


class Ctx:
    def __init__(self):
        self.macros, self.used, self.expansions = {}, set(), 0


class Stream:
    """Tokens with macro expansion on demand. Arguments are read unexpanded, as TeX does."""

    def __init__(self, toks, ctx):
        self.stack, self.ctx = list(reversed(toks)), ctx

    def pop(self):
        return self.stack.pop() if self.stack else None

    def peek(self):
        return self.stack[-1] if self.stack else None

    def push(self, toks):
        self.stack.extend(reversed(toks))

    def next(self, expand=True):
        while True:
            t = self.pop()
            if not expand or t is None or t.kind != "cs" or t.val not in self.ctx.macros:
                return t
            nargs, body = self.ctx.macros[t.val]
            self.ctx.used.add(t.val)
            self.ctx.expansions += 1
            if self.ctx.expansions > 200000:
                raise TexError(f"{t.where}: runaway expansion of \\{t.val}")
            args = [self.arg(t) for _ in range(nargs)]
            self.push([a for b in body for a in (args[b.val - 1] if b.kind == "param"
                                                  else [Tok(b.kind, b.val, t.where)])])

    def take_char(self, ch):
        t = self.peek()
        if t is not None and t.kind == "char" and t.val == ch:
            self.pop()
            return True
        return False

    def skip_space(self):
        while self.peek() is not None and self.peek().kind == "space":
            self.pop()

    def arg(self, owner=None):
        """One argument: a {group} without its braces, or a single token."""
        self.skip_space()
        t = self.pop()
        if t is None:
            raise TexError(f"{owner.where if owner else '?'}: missing argument")
        if t.kind != "bg":
            return [t]
        depth, out = 1, []
        while True:
            u = self.pop()
            if u is None:
                raise TexError(f"{t.where}: unbalanced {{")
            depth += (u.kind == "bg") - (u.kind == "eg")
            if depth == 0:
                return out
            out.append(u)

    def optional(self):
        self.skip_space()
        if not self.take_char("["):
            return None
        depth, out = 0, []
        while True:
            u = self.pop()
            if u is None:
                raise TexError("unclosed [")
            depth += (u.kind == "bg") - (u.kind == "eg")
            if depth == 0 and u.kind == "char" and u.val == "]":
                return out
            out.append(u)

    def raw(self, owner=None):
        return "".join(tok_text(t) for t in self.arg(owner)).strip()


def read_definition(s, t):
    """After \\newcommand: {\\name}[n]{body}. Returns (name, (nargs, body))."""
    s.skip_space()
    s.take_char("*")
    name = [x for x in s.arg(t) if x.kind != "space"]
    if len(name) != 1 or name[0].kind != "cs":
        raise TexError(f"{t.where}: cannot read the macro name")
    opt = s.optional()
    nargs = int("".join(tok_text(x) for x in opt)) if opt else 0
    if s.optional() is not None:
        raise TexError(f"{t.where}: macros with an optional argument are not supported")
    return name[0].val, (nargs, s.arg(t))


def resolve_tex(name):
    for cand in (name, name + ".tex"):
        path = os.path.normpath(os.path.join(HERE, cand))
        if os.path.isfile(path):
            return path
    return None


def scan_definitions(toks, ctx, seen):
    """Collect \\newcommand definitions from a preamble and the files it \\inputs."""
    s = Stream(toks, ctx)
    while (t := s.next(expand=False)) is not None:
        if t.kind != "cs":
            continue
        if t.val in ("newcommand", "renewcommand", "providecommand"):
            name, mac = read_definition(s, t)
            if t.val != "providecommand" or name not in ctx.macros:
                ctx.macros[name] = mac
        elif t.val == "input":
            name = s.raw(t)
            path = resolve_tex(name)
            if path is None:
                raise TexError(f"{t.where}: cannot find \\input{{{name}}}")
            if path not in seen:
                seen.add(path)
                scan_definitions(tokenize(read(path), os.path.relpath(path, HERE)), ctx, seen)


def read_aux(path):
    txt = read(path)
    labels = dict(re.findall(r"\\newlabel\{([^}]*)\}\{\{([^}]*)\}", txt))
    bibcite = {k: int(v) for k, v in re.findall(r"\\bibcite\{([^}]*)\}\{(\d+)\}", txt)}
    return labels, bibcite


def roman(n):
    out = ""
    for v, r in ((1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"),
                 (50, "L"), (40, "XL"), (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")):
        while n >= v:
            out, n = out + r, n - v
    return out


# ------------------------------------------------------------------ segments
BASE = dict(italic=False, bold=False, mono=False, sup=False, sub=False, color=None, font=None)
FMT_KEYS = tuple(BASE)
BREAK = "\n"


def emit(out, text, fmt):
    if not text:
        return
    key = tuple(fmt[k] for k in FMT_KEYS)
    if out and out[-1][1] == key and text != BREAK and out[-1][0] != BREAK:
        out[-1][0] += text
    else:
        out.append([text, key])


def tidy(segs):
    """Collapse spaces across runs and trim the ends, as TeX does."""
    out, last_space = [], True
    for text, key in segs:
        if text == BREAK:
            out.append([text, key])
            last_space = True
            continue
        text = re.sub(" +", " ", text)
        if last_space and text.startswith(" "):
            text = text[1:]
        if text:
            out.append([text, key])
            last_space = text.endswith(" ")
    while out and out[-1][0] != BREAK and out[-1][0].endswith(" "):
        out[-1][0] = out[-1][0].rstrip(" ")
        if not out[-1][0]:
            out.pop()
    return out


def plain(segs):
    return "".join(t for t, _ in segs if t != BREAK)


# ------------------------------------------------------------------ math
GREEK = dict(zip("alpha beta gamma delta epsilon zeta eta theta iota kappa lambda mu nu xi pi rho "
                 "sigma tau upsilon phi chi psi omega".split(),
                 "\u03b1\u03b2\u03b3\u03b4\u03f5\u03b6\u03b7\u03b8\u03b9\u03ba\u03bb\u03bc\u03bd"
                 "\u03be\u03c0\u03c1\u03c3\u03c4\u03c5\u03d5\u03c7\u03c8\u03c9", strict=True))
GREEK.update(varepsilon="\u03b5", varphi="\u03c6", vartheta="\u03d1", varrho="\u03f1")
GREEK_UP = dict(zip("Gamma Delta Theta Lambda Xi Pi Sigma Upsilon Phi Psi Omega".split(),
                    "\u0393\u0394\u0398\u039b\u039e\u03a0\u03a3\u03a5\u03a6\u03a8\u03a9", strict=True))
MATH_SYM = {  # name: (class, text)
    "in": ("rel", "\u2208"), "notin": ("rel", "\u2209"), "le": ("rel", "\u2264"),
    "leq": ("rel", "\u2264"), "ge": ("rel", "\u2265"), "geq": ("rel", "\u2265"),
    "ne": ("rel", "\u2260"), "neq": ("rel", "\u2260"), "approx": ("rel", "\u2248"),
    "sim": ("rel", "\u223c"), "equiv": ("rel", "\u2261"), "to": ("rel", "\u2192"),
    "rightarrow": ("rel", "\u2192"), "subset": ("rel", "\u2282"), "subseteq": ("rel", "\u2286"),
    "times": ("bin", "\u00d7"), "cdot": ("bin", "\u00b7"), "pm": ("bin", "\u00b1"),
    "cup": ("bin", "\u222a"), "cap": ("bin", "\u2229"), "bmod": ("bin", "mod"),
    "infty": ("ord", "\u221e"), "ldots": ("ord", "\u2026"), "dots": ("ord", "\u2026"),
    "cdots": ("ord", "\u22ef"), "partial": ("ord", "\u2202"), "ell": ("ord", "\u2113"),
    "%": ("ord", "%"), "_": ("ord", "_"), "#": ("ord", "#"),
    "{": ("open", "{"), "}": ("close", "}"), "lvert": ("open", "|"), "rvert": ("close", "|"),
    ",": ("space", THIN), ":": ("space", MEDSP), ";": ("space", " "), " ": ("space", " "),
    "!": ("space", ""), "quad": ("space", "\u2003"), "qquad": ("space", "\u2003\u2003"),
}
MATH_FUNCS = set("log ln exp sin cos tan max min sup inf lim arg det gcd Pr".split())
MATH_CHAR = {"+": "bin", "-": "bin", "*": "bin", "=": "rel", "<": "rel", ">": "rel", ":": "rel",
             "(": "open", "[": "open", ")": "close", "]": "close", ",": "punct", ";": "punct",
             "/": "ord", "|": "ord", "!": "ord", "'": "ord", ".": "punct"}


def A(cls, text, style="rm"):
    return ("atom", cls, text, style)


class MathReader:
    def __init__(self, reader):
        self.r = reader

    def parse(self, s, style="it"):
        nodes = []
        while (t := s.next()) is not None:
            if t.kind in ("space", "par"):
                continue
            if t.kind == "tie":
                nodes.append(A("space", " "))
            elif t.kind == "bg":
                s.push([t])
                nodes.append(("group", self.parse(Stream(s.arg(t), s.ctx), style)))
            elif t.kind == "char":
                self.char(t, s, nodes, style)
            elif t.kind == "cs":
                self.command(t, s, nodes, style)
            else:
                raise TexError(f"{t.where}: unexpected {t.kind} in math")
        return nodes

    def char(self, t, s, nodes, style):
        c = t.val
        if c in "_^":
            nt = s.next()
            if nt is None:
                raise TexError(f"{t.where}: {c} without argument")
            arg = []
            if nt.kind == "cs":
                self.command(nt, s, arg, style)
            else:
                s.push([nt])
                arg = self.parse(Stream(s.arg(t), s.ctx), style)
            base = nodes.pop() if nodes else A("ord", "")
            if base[0] != "script":
                base = ["script", base, None, None]
            slot = 2 if c == "_" else 3
            if base[slot] is not None:
                raise TexError(f"{t.where}: double {c}")
            base[slot] = arg
            nodes.append(base)
        elif c.isdigit() or (c == "." and s.peek() is not None and s.peek().kind == "char"
                             and s.peek().val.isdigit()):
            num = c
            while s.peek() is not None and s.peek().kind == "char" and (
                    s.peek().val.isdigit() or s.peek().val == "."):
                num += s.pop().val
            nodes.append(A("ord", num))
        elif c.isalpha():
            nodes.append(A("ord", c, style))
        elif c in MATH_CHAR:
            nodes.append(A(MATH_CHAR[c], {"-": "\u2212", "*": "\u2217", "'": "\u2032"}.get(c, c)))
        else:
            raise TexError(f"{t.where}: unsupported character {c!r} in math")

    def command(self, t, s, nodes, style):
        n = t.val
        if n in GREEK:
            nodes.append(A("ord", GREEK[n], "it"))
        elif n in GREEK_UP:
            nodes.append(A("ord", GREEK_UP[n]))
        elif n in MATH_SYM:
            nodes.append(A(*MATH_SYM[n]))
        elif n in MATH_FUNCS:
            nodes.append(A("func", n))
        elif n in ("mathrm", "operatorname", "mathsf", "mathtt", "mathbf", "mathit"):
            st = "it" if n == "mathit" else "rm"
            nodes.append(("group", self.parse(Stream(s.arg(t), s.ctx), st)))
        elif n in ("text", "textrm", "mbox"):
            nodes.append(A("ord", plain(self.r.segs(s.arg(t))), "text"))
        elif n == "frac":
            num = self.parse(Stream(s.arg(t), s.ctx), style)
            nodes.append(("frac", num, self.parse(Stream(s.arg(t), s.ctx), style)))
        elif n in ("left", "right", "big", "Big", "bigl", "bigr", "displaystyle"):
            s.take_char(".")
        else:
            raise TexError(f"{t.where}: unsupported math command \\{n}")


def math_segments(nodes, fmt, out, prev=None):
    """Inline math as text runs: italic letters, upright digits, subscripts via vertAlign."""
    def put(text, f):  # only the glyphs Times lacks go to the math font
        for chunk in re.split(f"([{''.join(sorted(NOT_IN_TIMES))}])", text):
            emit(out, chunk, dict(f, font=MATH_FONT) if chunk in NOT_IN_TIMES else f)

    for node in nodes:
        if node[0] == "group":
            prev = math_segments(node[1], fmt, out, prev)
            continue
        if node[0] == "script":
            prev = math_segments([node[1]], fmt, out, prev)
            for slot, key in ((2, "sub"), (3, "sup")):
                if node[slot] is not None:
                    math_segments(node[slot], dict(fmt, **{key: True}), out)
            continue
        if node[0] == "frac":
            math_segments(node[1], fmt, out)
            put("/", fmt)
            prev = math_segments(node[2], fmt, out)
            continue
        _, cls, text, style = node
        if cls == "ord":
            put(text, dict(fmt, italic=style == "it"))
        elif cls == "bin" and text == "mod":
            put(NBSP + "mod" + NBSP, fmt)
        elif cls == "bin" and prev in (None, "bin", "rel", "open", "punct"):
            put(text, fmt)
        elif cls in ("bin", "rel"):
            put(NBSP + text + NBSP, fmt)
        elif cls == "punct" and text == ",":
            put("," + THIN, fmt)
        elif cls == "func":
            put(text + THIN, fmt)
        else:
            put(text, fmt)
        prev = cls
    return prev


def m_el(tag, parent=None):
    el = OxmlElement(tag)
    if parent is not None:
        parent.append(el)
    return el


def m_run(parent, text, upright=False):
    r = m_el("m:r", parent)
    if upright:
        m_el("m:sty", m_el("m:rPr", r)).set(qn("m:val"), "p")
    rf = m_el("w:rFonts", m_el("w:rPr", r))
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), MATH_FONT)
    t = m_el("m:t", r)
    t.text = text
    t.set(qn("xml:space"), "preserve")


def omml(nodes, parent):
    """Display math as a native Word equation (OMML)."""
    for node in nodes:
        if node[0] == "group":
            omml(node[1], parent)
        elif node[0] == "script":
            tag = {(True, False): "m:sSub", (False, True): "m:sSup"}.get(
                (node[2] is not None, node[3] is not None), "m:sSubSup")
            el = m_el(tag, parent)
            omml([node[1]], m_el("m:e", el))
            if node[2] is not None:
                omml(node[2], m_el("m:sub", el))
            if node[3] is not None:
                omml(node[3], m_el("m:sup", el))
        elif node[0] == "frac":
            el = m_el("m:f", parent)
            omml(node[1], m_el("m:num", el))
            omml(node[2], m_el("m:den", el))
        else:
            _, cls, text, style = node
            if cls == "bin" and text == "mod":
                m_run(parent, MEDSP + "mod" + MEDSP, upright=True)
            else:
                m_run(parent, text, upright=(cls == "ord" and style != "it") or cls == "func")


# ------------------------------------------------------------------ reader
FMT_CMDS = {
    "emph": lambda f: f.update(italic=not f["italic"]), "textit": lambda f: f.update(italic=True),
    "textsl": lambda f: f.update(italic=True), "textbf": lambda f: f.update(bold=True),
    "texttt": lambda f: f.update(mono=True), "textup": lambda f: f.update(italic=False),
    "textrm": lambda f: f.update(mono=False), "textnormal": lambda f: f.update(BASE),
    "textsuperscript": lambda f: f.update(sup=True), "textsubscript": lambda f: f.update(sub=True),
    "IEEEauthorblockN": lambda f: None, "mbox": lambda f: None, "hbox": lambda f: None,
}
DECLS = {
    "bfseries": lambda f: f.update(bold=True), "itshape": lambda f: f.update(italic=True),
    "em": lambda f: f.update(italic=not f["italic"]), "ttfamily": lambda f: f.update(mono=True),
    "rmfamily": lambda f: f.update(mono=False), "upshape": lambda f: f.update(italic=False),
    "normalfont": lambda f: f.update(BASE),
}
IGNORED_DECLS = set("selectfont normalsize small footnotesize scriptsize tiny large Large LARGE "
                    "huge Huge centering raggedright raggedleft relax protect xspace "
                    "BIBentryALTinterwordspacing BIBentrySTDinterwordspacing".split())
SYMBOLS = {
    "%": "%", "_": "_", "&": "&", "#": "#", "$": "$", "{": "{", "}": "}", " ": " ", ",": THIN,
    ";": " ", ":": " ", "!": "", "/": "", "-": "", "@": "", "ldots": "\u2026", "dots": "\u2026",
    "textellipsis": "\u2026", "LaTeX": "LaTeX", "TeX": "TeX", "textquoteright": "\u2019",
    "textquoteleft": "\u2018", "textquotedblleft": "\u201c", "textquotedblright": "\u201d",
    "textasciitilde": "~", "textbar": "|", "textless": "<", "textgreater": ">",
    "textdegree": "\u00b0", "textmu": "\u00b5", "texttimes": "\u00d7", "ss": "\u00df",
    "o": "\u00f8", "O": "\u00d8", "ae": "\u00e6", "aa": "\u00e5", "l": "\u0142", "i": "\u0131",
    "newblock": " ", "quad": " ", "qquad": " ", "hfill": " ", "allowbreak": ZWSP,
}
ACCENTS = {'"': "\u0308", "'": "\u0301", "`": "\u0300", "^": "\u0302", "~": "\u0303",
           "=": "\u0304", ".": "\u0307", "u": "\u0306", "v": "\u030c", "H": "\u030b",
           "c": "\u0327", "k": "\u0328", "r": "\u030a"}
# Page layout commands with (optional, mandatory) argument counts: the template does this job.
LAYOUT = {"pagenumbering": (0, 1), "setcounter": (0, 2), "thispagestyle": (0, 1),
          "pagestyle": (0, 1), "chead": (0, 1), "lhead": (0, 1), "rhead": (0, 1),
          "cfoot": (0, 1), "lfoot": (0, 1), "rfoot": (0, 1), "fancyhead": (1, 1),
          "fancyfoot": (1, 1), "fancyhf": (0, 1), "bibliographystyle": (0, 1),
          "IEEEpeerreviewmaketitle": (0, 0), "balance": (0, 0), "vspace": (0, 1),
          "vfill": (0, 0), "IEEEtriggeratref": (0, 1), "bstctlcite": (0, 1)}
TWIPS_PER = {"cm": 1440 / 2.54, "mm": 144 / 2.54, "in": 1440, "pt": 1440 / 72.27, "bp": 20,
             "\\columnwidth": COLUMN_TWIPS, "\\linewidth": COLUMN_TWIPS}


def parse_colspec(spec, t):
    """A tabular column spec as [(alignment, width in twips or None)]: l c r p{..} and the
    array package's >{..} <{..} @{..}. Vertical rules are dropped."""
    s, cols, pre = Stream(spec, None), [], ""
    while (x := s.next(expand=False)) is not None:
        c = x.val if x.kind == "char" else None
        if x.kind == "space" or c == "|":
            continue
        if c in ("@", "!", "<"):
            s.arg(t)
        elif c == ">":
            pre = "".join(tok_text(y) for y in s.arg(t))
        elif c in ("l", "c", "r"):
            cols.append(({"l": "left", "c": "center", "r": "right"}[c], None))
            pre = ""
        elif c in ("p", "m", "b"):
            size = "".join(tok_text(y) for y in s.arg(t)).replace(" ", "")
            m = re.fullmatch(r"([\d.]*)(cm|mm|in|pt|bp|\\columnwidth|\\linewidth)", size)
            if not m:
                raise TexError(f"{t.where}: cannot read the column width {size}")
            align = ("center" if "\\centering" in pre else "right" if "\\raggedleft" in pre
                     else "left" if "\\raggedright" in pre else "justify")
            cols.append((align, round(float(m.group(1) or 1) * TWIPS_PER[m.group(2)])))
            pre = ""
        else:
            raise TexError(f"{t.where}: unsupported column type {tok_text(x)}")
    return cols


class Reader:
    def __init__(self, ctx, aux_labels, bibcite, bbl_path):
        self.ctx, self.aux, self.bibcite, self.bbl_path = ctx, aux_labels, bibcite, bbl_path
        self.blocks, self.para, self.indent, self.fmt = [], [], True, dict(BASE)
        self.title = self.author = None
        self.thanks, self.labels, self.dashes = [], {}, 0
        self.count = {"section": 0, "subsection": 0, "equation": 0, "figure": 0, "table": 0}
        self.label_value = None

    # -- inline content
    def segs(self, toks, fmt=None):
        out = []
        self.inline(Stream(toks, self.ctx), dict(fmt or BASE), out)
        return tidy(out)

    def inline(self, s, fmt, out):
        while (t := s.next()) is not None:
            self.inline_tok(t, s, fmt, out)

    def inline_tok(self, t, s, fmt, out):
        if t.kind == "char":
            self.text_char(t, s, fmt, out)
        elif t.kind in ("space", "par"):
            emit(out, " ", fmt)
        elif t.kind == "tie":
            emit(out, NBSP, fmt)
        elif t.kind == "bg":
            s.push([t])
            self.inline(Stream(s.arg(t), self.ctx), dict(fmt), out)
        elif t.kind == "math":
            toks = []
            while (u := s.next(expand=False)) is None or u.kind != "math":
                if u is None:
                    raise TexError(f"{t.where}: unclosed $")
                toks.append(u)
            nodes = MathReader(self).parse(Stream(toks, self.ctx))
            math_segments(nodes, dict(fmt, italic=False, mono=False), out)
        elif t.kind == "cs":
            self.inline_cs(t, s, fmt, out)
        else:
            raise TexError(f"{t.where}: unexpected {t.kind} {t.val!r}")

    def text_char(self, t, s, fmt, out):
        c = t.val
        if fmt["mono"]:
            emit(out, c, fmt)
        elif c == "-" and s.take_char("-"):
            s.take_char("-")
            self.dashes += 1
            emit(out, "-", fmt)
        elif c == "`":
            emit(out, "\u201c" if s.take_char("`") else "\u2018", fmt)
        elif c == "'":
            emit(out, "\u201d" if s.take_char("'") else "\u2019", fmt)
        else:
            emit(out, c, fmt)

    def inline_cs(self, t, s, fmt, out):
        n = t.val
        if n in SYMBOLS:
            emit(out, SYMBOLS[n], fmt)
        elif n in FMT_CMDS:
            f = dict(fmt)
            FMT_CMDS[n](f)
            self.inline(Stream(s.arg(t), self.ctx), f, out)
        elif n in DECLS:
            DECLS[n](fmt)
        elif n in IGNORED_DECLS:
            pass
        elif n in ACCENTS:
            base = plain(self.segs(s.arg(t))) or " "
            emit(out, unicodedata.normalize("NFC", base[0] + ACCENTS[n] + base[1:]), fmt)
        elif n in ("textendash", "textemdash"):
            self.dashes += 1
            emit(out, "-", fmt)
        elif n == "\\":
            s.optional()
            emit(out, BREAK, fmt)
        elif hasattr(self, "cs_" + n):
            getattr(self, "cs_" + n)(t, s, fmt, out)
        else:
            raise TexError(f"{t.where}: unsupported command \\{n}")

    def cs_cite(self, t, s, fmt, out):
        if s.optional() is not None:
            raise TexError(f"{t.where}: \\cite[...] is not supported")
        nums = []
        for key in (k.strip() for k in s.raw(t).split(",")):
            if key not in self.bibcite:
                raise TexError(f"{t.where}: citation {key} is not in main.aux")
            nums.append(self.bibcite[key])
        nums, parts, i = sorted(set(nums)), [], 0
        while i < len(nums):  # the cite package compresses three or more into a range
            j = i
            while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                j += 1
            if j - i >= 2:
                parts.append(f"[{nums[i]}]-[{nums[j]}]")
                i = j + 1
            else:
                parts.append(f"[{nums[i]}]")
                i += 1
        emit(out, ", ".join(parts), fmt)

    def ref_value(self, t, key):
        if key not in self.aux:
            raise TexError(f"{t.where}: \\ref{{{key}}} is not in main.aux")
        return self.aux[key]

    def cs_WordTitleBreak(self, t, s, fmt, out):  # empty in LaTeX; a line break in the Word title
        emit(out, BREAK, fmt)

    def cs_ref(self, t, s, fmt, out):
        emit(out, self.ref_value(t, s.raw(t)), fmt)

    def cs_eqref(self, t, s, fmt, out):
        emit(out, f"({self.ref_value(t, s.raw(t))})", fmt)

    def cs_label(self, t, s, fmt, out):
        key = s.raw(t)
        if self.label_value is None:
            raise TexError(f"{t.where}: \\label{{{key}}} outside a numbered item")
        self.labels[key] = (self.label_value, t.where)

    def cs_href(self, t, s, fmt, out):
        s.raw(t)
        self.inline(Stream(s.arg(t), self.ctx), dict(fmt), out)

    def cs_url(self, t, s, fmt, out):  # break after / and _ as \UrlBreaks does
        emit(out, re.sub(r"([/_])", "\\1" + ZWSP, s.raw(t)), dict(fmt))

    def cs_textcolor(self, t, s, fmt, out):
        color = s.raw(t)
        if color not in COLORS:
            raise TexError(f"{t.where}: unknown color {color}")
        self.inline(Stream(s.arg(t), self.ctx), dict(fmt, color=COLORS[color]), out)

    def cs_thanks(self, t, s, fmt, out):
        self.thanks.append(s.arg(t))

    def cs_fontsize(self, t, s, fmt, out):
        s.arg(t)
        s.arg(t)

    def cs_hskip(self, t, s, fmt, out):  # bbl glue: \hskip 1em plus 0.5em minus 0.4em\relax
        while s.peek() is not None and s.peek().kind in ("char", "space") and (
                s.peek().kind == "space" or s.peek().val.isalnum() or s.peek().val in ".-+"):
            s.pop()  # the glue specification only, so a missing \relax cannot eat text
        if s.peek() is not None and s.peek().kind == "cs" and s.peek().val == "relax":
            s.pop()
        emit(out, " ", fmt)

    def cs_bibinfo(self, t, s, fmt, out):
        s.raw(t)
        self.inline(Stream(s.arg(t), self.ctx), dict(fmt), out)

    cs_BIBforeignlanguage = cs_bibinfo

    def cs_ensuremath(self, t, s, fmt, out):
        nodes = MathReader(self).parse(Stream(s.arg(t), self.ctx))
        math_segments(nodes, dict(fmt, italic=False, mono=False), out)

    def cs_noindent(self, t, s, fmt, out):
        self.indent = False

    def cs_newcommand(self, t, s, fmt, out):
        name, mac = read_definition(s, t)
        self.ctx.macros[name] = mac

    cs_renewcommand = cs_providecommand = cs_newcommand

    # -- blocks
    def flush(self):
        segs = tidy(self.para)
        if segs:
            self.blocks.append(("para", segs, self.indent))
        self.para, self.indent, self.fmt = [], True, dict(BASE)

    def body(self, s):
        while (t := s.next()) is not None:
            if t.kind == "par":
                self.flush()
            elif t.kind == "cs" and t.val == "begin":
                self.environment(s.raw(t), t, s)
            elif t.kind == "cs" and t.val == "end":
                env = s.raw(t)
                if env != "document":
                    raise TexError(f"{t.where}: \\end{{{env}}} without \\begin")
                self.flush()
                return
            elif t.kind == "cs" and hasattr(self, "blk_" + t.val):
                self.flush()
                getattr(self, "blk_" + t.val)(t, s)
            elif t.kind == "cs" and t.val in LAYOUT:
                opt, mand = LAYOUT[t.val]
                for _ in range(opt):
                    s.optional()
                for _ in range(mand):
                    s.arg(t)
            else:
                self.inline_tok(t, s, self.fmt, self.para)
        raise TexError("main.tex: \\end{document} not found")

    def blk_title(self, t, s):
        self.title = s.arg(t)

    def blk_author(self, t, s):
        self.author = s.arg(t)

    def blk_maketitle(self, t, s):
        if self.title is None or self.author is None:
            raise TexError(f"{t.where}: \\maketitle before \\title and \\author")
        title, author = self.segs(self.title), self.segs(self.author)
        self.blocks.append(("title", title, author, [self.segs(x) for x in self.thanks]))

    def heading(self, t, s, level):
        starred = s.take_char("*")
        s.optional()
        if starred:
            self.label_value = None
        elif level == 1:
            self.count["section"] += 1
            self.count["subsection"] = 0
            self.label_value = roman(self.count["section"])
        else:
            self.count["subsection"] += 1
            self.label_value = f"{roman(self.count['section'])}-{chr(64 + self.count['subsection'])}"
        self.blocks.append(("heading", 5 if starred else level, self.segs(s.arg(t))))

    def blk_section(self, t, s):
        self.heading(t, s, 1)

    def blk_subsection(self, t, s):
        self.heading(t, s, 2)

    def blk_bibliography(self, t, s):
        s.raw(t)
        self.blocks.append(("references", self.bibliography()))

    def env_tokens(self, env, t, s):
        out, depth = [], 1
        while True:
            u = s.next(expand=False)
            if u is None:
                raise TexError(f"{t.where}: \\begin{{{env}}} is not closed")
            if u.kind == "cs" and u.val in ("begin", "end"):
                arg = s.arg(u)
                if "".join(tok_text(x) for x in arg) == env:
                    depth += 1 if u.val == "begin" else -1
                    if depth == 0:
                        return out
                out += [u, Tok("bg", "{", u.where)] + arg + [Tok("eg", "}", u.where)]
            else:
                out.append(u)

    def environment(self, env, t, s):
        body = self.env_tokens(env, t, s)
        if env == "abstract":
            self.flush()
            paras, cur = [], []
            for u in body + [Tok("par", None, t.where)]:
                if u.kind == "par":
                    if cur:
                        paras.append(self.segs(cur))
                    cur = []
                else:
                    cur.append(u)
            self.blocks.append(("abstract", [p for p in paras if p]))
        elif env == "IEEEkeywords":
            self.flush()
            self.blocks.append(("keywords", self.segs(body)))
        elif env in ("equation", "equation*"):
            continues = bool(tidy(self.para))
            self.flush()
            number = None
            if env == "equation":
                self.count["equation"] += 1
                number = str(self.count["equation"])
            self.label_value = number
            es, rest = Stream(body, self.ctx), []
            while (u := es.next(expand=False)) is not None:
                if u.kind == "cs" and u.val == "label":
                    self.cs_label(u, es, None, None)
                else:
                    rest.append(u)
            self.blocks.append(("equation", MathReader(self).parse(Stream(rest, self.ctx)), number))
            self.indent = not continues
        elif env in ("figure", "figure*"):
            self.figure(body, t)
        elif env == "table":
            self.table(body, t)
        else:
            raise TexError(f"{t.where}: unsupported environment {env}")

    def table(self, body, t):
        continues = bool(tidy(self.para))
        self.flush()
        self.count["table"] += 1
        self.label_value = roman(self.count["table"])  # IEEEtran numbers tables I, II, ...
        ts, caption, grid = Stream(body, self.ctx), None, None
        ts.optional()
        while (u := ts.next()) is not None:
            if u.kind in ("space", "par") or (u.kind == "cs" and u.val in IGNORED_DECLS):
                continue
            if u.kind == "cs" and u.val == "caption":
                ts.optional()
                caption = self.segs(ts.arg(u))
            elif u.kind == "cs" and u.val == "label":
                self.cs_label(u, ts, None, None)
            elif u.kind == "cs" and u.val in ("renewcommand", "newcommand"):
                read_definition(ts, u)  # \arraystretch and the like: layout only
            elif u.kind == "cs" and u.val == "begin" and ts.raw(u) == "tabular":
                grid = self.tabular(self.env_tokens("tabular", u, ts), u)
            else:
                raise TexError(f"{u.where}: unsupported in table: {tok_text(u)}")
        if caption is None or grid is None:
            raise TexError(f"{t.where}: a table needs \\caption{{...}} and a tabular")
        self.blocks.append(("table", caption, *grid))
        self.indent = not continues

    def tabular(self, toks, t):
        """Rows split at \\\\ and cells at &; \\hline positions become cell borders."""
        s = Stream(toks, self.ctx)
        s.optional()
        cols = parse_colspec(s.arg(t), t)
        rows, row, cell, rules, depth = [], [], [], set(), 0
        while (x := s.next(expand=False)) is not None:
            depth += (x.kind == "bg") - (x.kind == "eg")
            if depth == 0 and x.kind == "cs" and x.val == "hline":
                rules.add(len(rows))
            elif depth == 0 and x.kind == "char" and x.val == "&":
                row, cell = row + [cell], []
            elif depth == 0 and x.kind == "cs" and x.val == "\\":
                s.optional()
                rows.append(row + [cell])
                row, cell = [], []
            elif x.kind == "cs" and x.val in ("multicolumn", "multirow", "cline"):
                raise TexError(f"{x.where}: \\{x.val} is not supported")
            else:
                cell.append(x)
        if row or any(x.kind not in ("space", "par") for x in cell):
            rows.append(row + [cell])
        for r in rows:
            if len(r) != len(cols):
                raise TexError(f"{t.where}: a row has {len(r)} cells for {len(cols)} columns")
        return cols, [[self.segs(c) for c in r] for r in rows], rules

    def figure(self, body, t):
        continues = bool(tidy(self.para))
        self.flush()
        self.count["figure"] += 1
        self.label_value = str(self.count["figure"])
        fs, source, caption = Stream(body, self.ctx), None, None
        fs.optional()
        while (u := fs.next()) is not None:
            if u.kind in ("space", "par") or (u.kind == "cs" and u.val in IGNORED_DECLS):
                continue
            if u.kind == "cs" and u.val == "input":
                source = fs.raw(u)
            elif u.kind == "cs" and u.val == "caption":
                fs.optional()
                caption = self.segs(fs.arg(u))
            elif u.kind == "cs" and u.val == "label":
                self.cs_label(u, fs, None, None)
            else:
                raise TexError(f"{u.where}: unsupported in figure: {tok_text(u)}")
        if source is None or caption is None:
            raise TexError(f"{t.where}: a figure needs \\input{{...}} and \\caption{{...}}")
        self.blocks.append(("figure", source, caption))
        self.indent = not continues

    def bibliography(self):
        s = Stream(tokenize(read(self.bbl_path), "main.bbl"), self.ctx)
        while (t := s.next(expand=False)) is not None and not (t.kind == "cs" and t.val == "bibitem"):
            pass
        items = []
        while t is not None and t.kind == "cs" and t.val == "bibitem":
            s.optional()
            key, body = s.raw(t), []
            while (t := s.next(expand=False)) is not None and not (
                    t.kind == "cs" and t.val in ("bibitem", "end")):
                body.append(t)
            items.append((key, self.segs(body)))
        order = sorted(self.bibcite, key=self.bibcite.get)
        if [k for k, _ in items] != order:
            raise TexError("main.bbl order differs from the \\bibcite numbers in main.aux")
        return [segs for _, segs in items]

    def check_labels(self):
        for key, (ours, where) in self.labels.items():
            if self.aux.get(key) != ours:
                raise TexError(f"{where}: {key} is {ours} here but {self.aux.get(key)} in main.aux")


# ------------------------------------------------------------------ LaTeX runs
def run(cmd, cwd, env=None, log=None):
    res = subprocess.run(cmd, cwd=cwd, env=env, capture_output=True, text=True,
                         encoding="utf-8", errors="replace")
    if log:
        with open(log, "w", encoding="utf-8") as f:
            f.write(res.stdout + res.stderr)
    if res.returncode != 0:
        sys.exit(f"failed ({res.returncode}): {' '.join(cmd)}; see {log or 'output'}")


def build_latex():
    os.makedirs(LATEX_OUT, exist_ok=True)
    env = dict(os.environ, BIBINPUTS=HERE + os.pathsep)  # bibtex runs inside the outdir
    run(["latexmk", "-g", "-pdf", "-interaction=nonstopmode", "-halt-on-error",
         "-outdir=" + os.path.relpath(LATEX_OUT, HERE), "main.tex"], HERE, env,
        os.path.join(BUILD, "latexmk.log"))


def build_figure(name, preamble, compile_tex):
    """Standalone build of one figure with main.tex's preamble, rendered to PNG."""
    os.makedirs(FIG_OUT, exist_ok=True)
    stem = os.path.splitext(os.path.basename(name))[0]
    pdf, png = os.path.join(FIG_OUT, stem + ".pdf"), os.path.join(FIG_OUT, stem + ".png")
    if compile_tex:
        for ext in (".aux", ".out"):  # start clean: a stale file can break the next run
            if os.path.exists(os.path.join(FIG_OUT, stem + ext)):
                os.remove(os.path.join(FIG_OUT, stem + ext))
        wrapper = os.path.join(FIG_OUT, stem + ".tex")
        with open(wrapper, "w", encoding="utf-8", newline="\n") as f:
            f.write(preamble + "\\usepackage[active,tightpage]{preview}\n"
                    "\\PreviewEnvironment{tikzpicture}\n\\setlength\\PreviewBorder{1pt}\n"
                    "\\begin{document}\n\\input{" + name + "}\n\\end{document}\n")
        run(["pdflatex", "-interaction=nonstopmode", "-halt-on-error",
             "-output-directory=" + os.path.relpath(FIG_OUT, HERE), os.path.relpath(wrapper, HERE)],
            HERE, None, os.path.join(FIG_OUT, stem + "-console.txt"))  # .out is hyperref's
        with fitz.open(pdf) as d:
            d[0].get_pixmap(dpi=FIG_DPI).save(png)
    with fitz.open(pdf) as d:
        width_pt, height_pt = d[0].rect.width, d[0].rect.height
    return png, width_pt, height_pt


# ------------------------------------------------------------------ Word output
def set_font(run_, name):
    rf = run_._element.get_or_add_rPr().get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), name)


def add_runs(p, segs, font=None):
    for text, key in segs:
        f = dict(zip(FMT_KEYS, key, strict=True))
        if text == BREAK:
            p.add_run().add_break()
            continue
        r = p.add_run(text)
        if f["italic"]:
            r.italic = True
        if f["bold"]:
            r.bold = True
        if f["sup"]:
            r.font.superscript = True
        if f["sub"]:
            r.font.subscript = True
        if f["color"]:
            r.font.color.rgb = RGBColor.from_string(f["color"])
        name = CODE_FONT if f["mono"] else (f["font"] or font)
        if name:
            set_font(r, name)


class WordWriter:
    def __init__(self, template):
        self.doc = docx.Document(template)
        body = self.doc.element.body
        sects = list(body.iter(qn("w:sectPr")))
        self.title_sect = next(x for x in sects if x.find(qn("w:titlePg")) is not None)
        self.body_sect = next(x for x in sects if (x.find(qn("w:cols")) is not None and
                                                   x.find(qn("w:cols")).get(qn("w:num")) == "2"))
        sponsor = next(p for p in body.iter(qn("w:p")) if p.find(qn("w:pPr")) is not None and
                       p.find(qn("w:pPr")).find(qn("w:pStyle")) is not None and
                       p.find(qn("w:pPr")).find(qn("w:pStyle")).get(qn("w:val")) == "sponsors")
        self.sponsor_ppr = [copy.deepcopy(sponsor.find(qn("w:pPr")).find(qn(x)))
                            for x in ("w:framePr", "w:ind")]
        self.title_sect, self.body_sect = copy.deepcopy(self.title_sect), copy.deepcopy(self.body_sect)
        for child in list(body):
            body.remove(child)
        body.append(self.body_sect)  # the document ends in the two-column body section
        # Section numbers are followed by a tab to 576 twips, which "VIII." overruns, leaving
        # no gap; a space after the number spaces every heading alike, as LaTeX does.
        for lvl in self.doc.part.numbering_part.element.iter(qn("w:lvl")):
            style = lvl.find(qn("w:pStyle"))
            if (style is not None and style.get(qn("w:val")) == "Heading1"
                    and lvl.find(qn("w:suff")) is None):
                suff = OxmlElement("w:suff")
                suff.set(qn("w:val"), "space")
                lvl.find(qn("w:lvlText")).addprevious(suff)

    def para(self, style, segs=(), font=None):
        p = self.doc.add_paragraph(style=self.doc.styles[style])
        add_runs(p, segs, font)
        return p

    def title(self, title, author, thanks):
        self.para("paper title", title)
        a = self.para("Author", author)
        a.paragraph_format.space_after = Pt(12)  # the template's affiliation lines would sit here
        a._p.get_or_add_pPr().append(self.title_sect)  # one-column title section, first-page header
        for segs in thanks:  # the \thanks note: the template's frame at the foot of column 1
            p = self.para("sponsors")
            for el in self.sponsor_ppr:
                p._p.get_or_add_pPr().append(copy.deepcopy(el))
            add_runs(p, segs)

    def labelled(self, style, label, segs):
        p = self.para(style)
        r = p.add_run(label + LABEL_SEP + " ")
        r.italic = True
        add_runs(p, segs)

    def equation(self, nodes, number):
        p = self.para("equation")
        stops = p.paragraph_format.tab_stops
        for pos in (2520, 5040):  # the style's stops assume a wider column
            stops.add_tab_stop(Twips(pos), WD_TAB_ALIGNMENT.CLEAR)
        stops.add_tab_stop(Twips(COLUMN_TWIPS // 2), WD_TAB_ALIGNMENT.CENTER)
        stops.add_tab_stop(Twips(COLUMN_TWIPS - 40), WD_TAB_ALIGNMENT.RIGHT)
        set_font(p.add_run("\t"), TEXT_FONT)  # the style's font is Symbol
        omml(nodes, m_el("m:oMath", p._p))
        set_font(p.add_run("\t" + (f"({number})" if number else "")), TEXT_FONT)

    def figure(self, png, width_pt, height_pt, caption):
        scale = min(1.0, COLUMN_TWIPS / 20 / width_pt)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.add_run().add_picture(png, width=Emu(int(width_pt * scale * 12700)),
                                height=Emu(int(height_pt * scale * 12700)))
        self.para("figure caption", caption)

    def table(self, caption, cols, rows, rules):
        cap = self.para("table head", caption)
        cap.paragraph_format.keep_with_next = True
        pad = 60  # twips each side; a p{} width in LaTeX excludes the padding, so add it back
        fixed = sum(w + 2 * pad for _, w in cols if w)
        free = [i for i, (_, w) in enumerate(cols) if not w]
        rest = max(COLUMN_TWIPS - fixed, 0) // len(free) if free else 0
        widths = [w + 2 * pad if w else rest for _, w in cols]
        scale = min(1.0, COLUMN_TWIPS / sum(widths))
        widths = [int(w * scale) for w in widths]
        t = self.doc.add_table(rows=len(rows), cols=len(cols))
        t.alignment, t.autofit = WD_TABLE_ALIGNMENT.CENTER, False
        for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths, strict=True):
            gc.set(qn("w:w"), str(w))
        mar = OxmlElement("w:tblCellMar")
        for side in ("left", "right"):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:w"), str(pad))
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        look = t._tbl.tblPr.find(qn("w:tblLook"))
        if look is not None:
            look.addprevious(mar)  # schema order: ... tblLayout, tblCellMar, tblLook
        else:
            t._tbl.tblPr.append(mar)
        align = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                 "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
        header = 1 in rules and len(rows) > 1  # a rule under the first row marks a header
        for i, row in enumerate(rows):
            tr_pr = t.rows[i]._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement("w:cantSplit"))
            for j, segs in enumerate(row):
                cell = t.cell(i, j)
                cell.width = Twips(widths[j])
                p = cell.paragraphs[0]
                p.style = self.doc.styles["table col head" if header and i == 0 else "table copy"]
                p.alignment = align[cols[j][0]]
                p.paragraph_format.keep_with_next = i < len(rows) - 1
                add_runs(p, segs)
                edges = (["top"] if i in rules else []) + (
                    ["bottom"] if i == len(rows) - 1 and len(rows) in rules else [])
                if edges:
                    borders = OxmlElement("w:tcBorders")
                    cell._tc.get_or_add_tcPr().append(borders)  # after tcW, as the schema wants
                for edge in edges:
                    b = OxmlElement(f"w:{edge}")
                    for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"),
                                 ("w:color", "000000")):
                        b.set(qn(k), v)
                    borders.append(b)
        gap = self.doc.add_paragraph()
        gap.paragraph_format.line_spacing = Pt(6)

    def references(self, items):
        self.para("heading 5", [["References", tuple(BASE.values())]])
        for segs in items:
            p = self.para("references", segs)
            p.paragraph_format.left_indent = Twips(354)
            p.paragraph_format.first_line_indent = Twips(-354)


def write_docx(reader, figures, keywords_name):
    w = WordWriter(TEMPLATE)
    for b in reader.blocks:
        kind = b[0]
        if kind == "title":
            w.title(*b[1:])
        elif kind == "abstract":
            for i, segs in enumerate(b[1]):
                if i == 0:
                    w.labelled("Abstract", "Abstract", segs)
                else:
                    w.para("Abstract", segs)
        elif kind == "keywords":
            w.labelled("Keywords", keywords_name, b[1])
        elif kind == "heading":
            w.para(f"heading {b[1]}", b[2])
        elif kind == "para":
            p = w.para("Body Text", b[1])
            if not b[2]:
                p.paragraph_format.first_line_indent = Twips(0)
        elif kind == "equation":
            w.equation(b[1], b[2])
        elif kind == "figure":
            w.figure(*figures[b[1]], b[2])
        elif kind == "table":
            w.table(*b[1:])
        elif kind == "references":
            w.references(b[1])
    title = next(plain(b[1]) for b in reader.blocks if b[0] == "title")
    sup = FMT_KEYS.index("sup")  # the author name without its affiliation mark
    author = next("".join(t for t, k in b[2] if t != BREAK and not k[sup])
                  for b in reader.blocks if b[0] == "title").strip()
    cp = w.doc.core_properties
    cp.title, cp.author, cp.last_modified_by, cp.revision = title, author, author, 1
    cp.keywords = next((plain(b[1]) for b in reader.blocks if b[0] == "keywords"), "")
    cp.subject = cp.comments = ""
    w.doc.save(OUTPUT)
    normalize_zip(OUTPUT, title)


def clean_app_properties(data, title):
    """docProps/app.xml comes from the template: its title part, document security flag and
    statistics describe the template, not this paper. Word recomputes the statistics on save."""
    s = data.decode("utf-8")
    s = re.sub(r"<DocSecurity>\d+</DocSecurity>", "<DocSecurity>0</DocSecurity>", s)
    s = re.sub(r"(<TitlesOfParts><vt:vector[^>]*><vt:lpstr>)[^<]*(</vt:lpstr>)",
               lambda m: m.group(1) + xml_escape(title) + m.group(2), s)
    s = re.sub(r"<(Pages|Words|Characters|Lines|Paragraphs|CharactersWithSpaces)>\d+</\1>", "", s)
    return s.encode("utf-8")


def normalize_zip(path, title):
    """python-docx stamps zip entries with the current time; fix them for identical bytes.
    The template's application properties are cleaned on the way (clean_app_properties)."""
    with zipfile.ZipFile(path) as z:
        items = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    items = [(n, clean_app_properties(d, title) if n == "docProps/app.xml" else d) for n, d in items]
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in items:
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type, info.external_attr = zipfile.ZIP_DEFLATED, 0o644 << 16
            z.writestr(info, data)
    os.replace(tmp, path)


# ------------------------------------------------------------------ checks
def docx_text(path):
    with zipfile.ZipFile(path) as z:
        pat = r"word/(document|header\d*|footer\d*|footnotes)\.xml"
        parts = [n for n in z.namelist() if re.match(pat, n)]
        return "".join(re.sub(r"<[^>]+>", "", z.read(n).decode("utf-8")) for n in parts)


def check_output(ctx, reader):
    text = docx_text(OUTPUT)
    problems = [f"{name} found in the output" for ch, name in FORBIDDEN.items() if ch in text]
    if "\\" in text:
        problems.append("a backslash survived into the output")
    for name in sorted(n for n in ctx.used if n.startswith("Wake")):
        value = "".join(tok_text(x) for x in ctx.macros[name][1])
        if value not in text:
            problems.append(f"\\{name} = {value} is not in the output")
    if problems:
        sys.exit("check failed:\n  " + "\n  ".join(problems))
    kinds = [b[0] for b in reader.blocks]
    print(f"wrote {OUTPUT}")
    print(f"  {kinds.count('heading')} headings, {kinds.count('para')} paragraphs, "
          f"{kinds.count('equation')} equations, {kinds.count('figure')} figures, "
          f"{kinds.count('table')} tables, "
          f"{sum(len(b[1]) for b in reader.blocks if b[0] == 'references')} references, "
          f"{len([n for n in ctx.used if n.startswith('Wake')])} result macros expanded")
    if reader.dashes:
        print(f"  note: {reader.dashes} LaTeX dash(es) (-- or ---) written as '-'")


def word_check():
    """Open the .docx in Word itself (Windows): page count in Word's layout, and a PDF."""
    pdf = os.path.join(BUILD, os.path.splitext(os.path.basename(OUTPUT))[0] + ".pdf")
    ps = (f"$ErrorActionPreference='Stop'; $w=New-Object -ComObject Word.Application; "
          f"$w.Visible=$false; $w.DisplayAlerts=0; try {{ "
          f"$d=$w.Documents.Open('{OUTPUT}', $false, $true, $false); $d.Repaginate(); "
          f"$n=$d.ComputeStatistics(2); $d.ExportAsFixedFormat('{pdf}', 17); $d.Close(0); "
          f"Write-Output $n }} finally {{ $w.Quit() }}")
    res = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", ps],
                         capture_output=True, text=True, timeout=600)
    if res.returncode != 0:
        sys.exit("Word check failed: " + res.stderr.strip())
    pages = int(res.stdout.strip().splitlines()[-1])
    print(f"  Word: {pages} pages (limit {MAX_PAGES}); PDF at {pdf}")
    if pages > MAX_PAGES:
        sys.exit(f"over the {MAX_PAGES}-page limit")


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--no-latex", action="store_true", help="reuse build-docx/ from the last run")
    ap.add_argument("--word-check", action="store_true", help="count pages in Word (Windows)")
    args = ap.parse_args()
    if not args.no_latex:
        build_latex()
    src = read(MAIN_TEX)
    toks = tokenize(src, "main.tex")
    begin = "".join(tok_text(t) for t in tokenize("\\begin{document}", "-"))
    start = next((i for i in range(len(toks))
                  if "".join(tok_text(t) for t in toks[i:i + 11]) == begin), None)
    if start is None:
        sys.exit("build_docx: main.tex has no \\begin{document}")
    ctx = Ctx()
    scan_definitions(toks[:start], ctx, set())
    ctx.macros.pop("WordTitleBreak", None)  # not expanded: Reader.cs_WordTitleBreak breaks the line
    labels, bibcite = read_aux(os.path.join(LATEX_OUT, "main.aux"))
    reader = Reader(ctx, labels, bibcite, os.path.join(LATEX_OUT, "main.bbl"))
    reader.body(Stream(toks[start + 11:], ctx))
    reader.check_labels()
    preamble = src[:src.index("\\begin{document}")]
    figures = {b[1]: build_figure(b[1], preamble, not args.no_latex)
               for b in reader.blocks if b[0] == "figure"}
    kw = ctx.macros.get("IEEEkeywordsname")
    write_docx(reader, figures, plain(reader.segs(kw[1])) if kw else "Index Terms")
    check_output(ctx, reader)
    avail = ctx.macros.get("PaperRepoAvailability")  # main.tex's switch for code availability
    if avail:
        print(f"  code availability: \"{plain(reader.segs(avail[1]))}\" (\\PaperRepoAvailability)")
    if args.word_check:
        word_check()


if __name__ == "__main__":
    try:
        main()
    except TexError as e:
        sys.exit(f"build_docx: {e}")
